"""Shared helpers for the pytest suite (API calls and reading the Excel workbook)."""
import io
import os
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "test-data" / "upload_test_data.xlsx"
API_URL = os.getenv("API_URL", "http://localhost:4000").rstrip("/")
UI_URL = os.getenv("UI_URL", "http://localhost:5173").rstrip("/")   # use http://localhost:3000 for Docker mode
TEST_DOMAIN = "testdata.example"                                    # only records with this domain are touched
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Employees that must already exist before the upload (for the "replace" checks)
EXISTING = [
    {"full_name": "Existing Candidate", "email": f"existing.candidate@{TEST_DOMAIN}",
     "department": "Sales", "position": "Rep", "salary": 400000},
    {"full_name": "Existing Mixed", "email": f"existing.mixed@{TEST_DOMAIN}",
     "department": "Support", "position": "Agent", "salary": 300000},
]


# ---------- Excel ----------
def load_cases(path=WORKBOOK):
    """Read the 'Expected' sheet -> list of dicts (one per test case)."""
    ws = load_workbook(path, data_only=True)["Expected"]
    headers = [c.value for c in ws[1]]
    cases = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(values):
            continue
        case = dict(zip(headers, values))
        case["row"] = int(case["row"])
        case["expected_error_contains"] = (case.get("expected_error_contains") or "").strip()
        cases.append(case)
    return cases


def load_candidate_rows(path=WORKBOOK):
    """Read the 'Candidates' sheet -> {excel_row_number: {header: value}}."""
    ws = load_workbook(path, data_only=True)["Candidates"]
    headers = [c.value for c in ws[1]]
    return {i: dict(zip(headers, v)) for i, v in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)}


def make_xlsx_bytes(rows):
    """Build a small in-memory .xlsx from a list of rows (first row = headers)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- API ----------
def upload(source, mode="validate", replace_existing=True, filename="upload.xlsx"):
    """POST an Excel file (path or bytes) to the import endpoint."""
    if isinstance(source, (str, Path)):
        filename, data = Path(source).name, Path(source).read_bytes()
    else:
        data = source
    return requests.post(
        f"{API_URL}/api/employees/import",
        params={"mode": mode, "replaceExisting": str(replace_existing).lower()},
        files={"file": (filename, data, XLSX_MIME)},
        timeout=30,
    )


def list_test_employees():
    resp = requests.get(f"{API_URL}/api/employees", params={"search": TEST_DOMAIN}, timeout=15)
    resp.raise_for_status()
    return [e for e in resp.json() if e["email"].lower().endswith("@" + TEST_DOMAIN)]


def find_by_email(email):
    return [e for e in list_test_employees() if e["email"].lower() == email.lower()]


def delete_test_employees():
    for e in list_test_employees():
        requests.delete(f"{API_URL}/api/employees/{e['id']}", timeout=15).raise_for_status()


def seed_existing():
    delete_test_employees()
    for e in EXISTING:
        resp = requests.post(f"{API_URL}/api/employees", json=e, timeout=15)
        assert resp.status_code == 201, f"could not create {e['email']}: {resp.text}"
