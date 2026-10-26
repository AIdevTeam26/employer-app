#!/usr/bin/env python3
"""Builds test-data/upload_test_data.xlsx

Sheet "Candidates": the file you upload in the website (one test case per row).
Sheet "Expected"  : what the application MUST answer for each row (read by pytest).
Sheet "About"     : short instructions.

Edit the CASES list below and run:   python tests/build_test_workbook.py
or edit the Excel file by hand (keep the "row" numbers in sync with the Candidates sheet).
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent.parent / "test-data" / "upload_test_data.xlsx"
HEADERS = ["full_name", "email", "department", "position", "salary", "hire_date", "case_id"]
D = "@testdata.example"


def case(case_id, check, description, status, error, full_name, email, department, position,
         salary=None, hire_date=None):
    return dict(case_id=case_id, check=check, description=description, status=status, error=error,
                cells=[full_name, email, department, position, salary, hire_date, case_id])


CASES = [
    # ---- valid rows -------------------------------------------------------------------------
    case("TC01", "valid", "All fields filled correctly", "new", "",
         "Asha Menon", "asha.menon" + D, "Engineering", "QA Engineer", 800000, date(2024, 1, 15)),
    case("TC02", "valid", "Optional salary and hire_date left empty", "new", "",
         "Ravi Shankar", "ravi.shankar" + D, "Finance", "Accountant"),
    # ---- required fields --------------------------------------------------------------------
    case("TC03", "required", "full_name is empty", "invalid", "full_name is required",
         None, "missing.name" + D, "HR", "Recruiter", 500000, date(2023, 6, 1)),
    case("TC04", "required", "email is empty", "invalid", "email is required",
         "Missing Email", None, "HR", "Recruiter"),
    case("TC05", "required", "department is empty", "invalid", "department is required",
         "Missing Dept", "missing.dept" + D, None, "Analyst"),
    case("TC06", "required", "position is empty", "invalid", "position is required",
         "Missing Position", "missing.pos" + D, "Sales", None),
    case("TC07", "required", "full_name contains only spaces", "invalid", "full_name is required",
         "   ", "blank.name" + D, "Sales", "Rep"),
    # ---- data types -------------------------------------------------------------------------
    case("TC08", "datatype", "salary is text instead of a number", "invalid", "salary must be a number",
         "Salary Text", "salary.text" + D, "Sales", "Rep", "abc"),
    case("TC09", "datatype", "salary is negative", "invalid", "salary cannot be negative",
         "Salary Negative", "salary.negative" + D, "Sales", "Rep", -5000),
    case("TC10", "datatype", "hire_date is not a date", "invalid", "hire_date must be a valid date",
         "Date Text", "date.text" + D, "Sales", "Rep", 400000, "not a date"),
    case("TC11", "datatype", "hire_date is an impossible date (31 Feb)", "invalid",
         "hire_date must be a valid date",
         "Date Impossible", "date.impossible" + D, "Sales", "Rep", 400000, "2024-02-31"),
    case("TC12", "datatype", "full_name is a number instead of text", "invalid", "full_name must be text",
         12345, "name.number" + D, "Sales", "Rep"),
    case("TC13", "datatype", "full_name is longer than 120 characters", "invalid",
         "full_name must be 120 characters or fewer",
         "A" * 130, "long.name" + D, "Sales", "Rep"),
    # ---- email format -----------------------------------------------------------------------
    case("TC14", "email", "email has no @ sign", "invalid", "email is not valid",
         "Email No At", "john.doe.testdata.example", "IT", "Admin"),
    case("TC15", "email", "email has no domain", "invalid", "email is not valid",
         "Email No Domain", "john@", "IT", "Admin"),
    case("TC16", "email", "email contains a space", "invalid", "email is not valid",
         "Email Space", "john doe" + D, "IT", "Admin"),
    case("TC17", "email", "email domain has no extension", "invalid", "email is not valid",
         "Email No Tld", "john@localhost", "IT", "Admin"),
    # ---- replace existing candidates (the two existing records are created by the pytest fixture)
    case("TC18", "replace", "email already exists -> existing record is replaced", "replace", "",
         "Existing Candidate", "existing.candidate" + D, "Operations", "Team Lead", 1000000, date(2022, 3, 1)),
    case("TC19", "replace", "email already exists in different letter case -> replaced, not duplicated",
         "replace", "",
         "Existing Mixed", "EXISTING.MIXED@TESTDATA.EXAMPLE", "Operations", "Analyst"),
    # ---- duplicate inside the file ----------------------------------------------------------
    case("TC20", "duplicate", "same email twice in the file (repeats TC01)", "invalid",
         "Duplicate email in file",
         "Asha Duplicate", "asha.menon" + D, "Engineering", "QA Engineer"),
]


def style_header(ws, widths):
    fill = PatternFill("solid", fgColor="0F6B57")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()

    ws = wb.active
    ws.title = "Candidates"
    ws.append(HEADERS)
    for c in CASES:
        ws.append(c["cells"])
        r = ws.max_row
        if isinstance(c["cells"][5], date):
            ws.cell(row=r, column=6).number_format = "yyyy-mm-dd"
    style_header(ws, [30, 34, 16, 20, 12, 14, 9])

    ex = wb.create_sheet("Expected")
    ex.append(["row", "case_id", "check", "description", "expected_status", "expected_error_contains"])
    for i, c in enumerate(CASES, start=2):          # Candidates data starts on Excel row 2
        ex.append([i, c["case_id"], c["check"], c["description"], c["status"], c["error"]])
    style_header(ex, [6, 9, 11, 62, 16, 44])

    about = wb.create_sheet("About")
    for line in [
        "Test data for the Excel upload feature.",
        "Candidates = the file you upload in the website. Each row is one test case (see case_id).",
        "Expected   = the result the application must give for each row. pytest reads this sheet.",
        "The 'case_id' column in Candidates is ignored by the application.",
        "Existing records needed by the 'replace' cases are created automatically by the pytest fixture:",
        "  existing.candidate@testdata.example (Sales) and existing.mixed@testdata.example (Support).",
        "Expected status values: new | replace | skip | invalid",
    ]:
        about.append([line])
    about.column_dimensions["A"].width = 110

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT} with {len(CASES)} cases")


if __name__ == "__main__":
    main()
